"""SHEFISH 合成回归数据；不包含真实客户、订单或商品。"""

from __future__ import annotations

import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook


def generate(path: Path) -> None:
    workbook = Workbook()
    db = workbook.active
    assert db is not None
    db.title = "DATA BASE"
    db.append(
        [
            "Item Number",
            "COMBO DESCRIPTION",
            "Category",
            "round value",
            "N/W",
            "G/W",
            "L",
            "W",
            "H",
            "SK/YM COMBO FOB 2026",
            "YM ROD FOB 2026",
            "SK REEL FOB 2026",
            "GS PTE COMBO FOB 2026",
            "GS PTE ROD FOB 2026",
            "GS PTE REEL FOB 2026",
            "EMAX PTE COMBO DDP 2026",
            "EMAX PTE ROD DDP 2026",
            "EMAX PTE REEL DDP 2026",
        ]
    )
    for item, category in (
        ("TEST-COMBO", "COMBO"),
        ("TEST-ROD", "Single Rod"),
        ("TEST-REEL", "REEL"),
    ):
        db.append(
            [item, f"Master {item}", category, 6, 2, 3, 100, 10, 20, 10, 4, 6, 12, 5, 7, 20, 8, 12]
        )
    po = workbook.create_sheet("PO RECORD 2026")
    po.append(
        [
            "CATEGORY",
            "PI#",
            "PO#",
            "QTY",
            "Item Number",
            "UPC",
            "COMBO DESCRIPTION",
            "INV#",
            "FACTORY DOC NO.",
            2601,
            2602,
            "FINAL CONFIRMED EX-FACTORY DATE",
            "ETA",
            "CTNS",
            "TOTAL CBM",
        ]
    )
    cp = workbook.create_sheet("New PO")
    cp.append(["PO#", "Qty/Ea", "Item Number", "UPC", "Description", "PO Issue Date"])
    for item, category, quantity, upc in (
        ("TEST-COMBO", "COMBO", 120, "000123456789"),
        ("TEST-ROD", "Single Rod", 60, "000123456790"),
        ("TEST-REEL", "REEL", 90, "000123456791"),
    ):
        po.append(
            [
                category,
                "TEST-PI-01",
                "TEST-PO-01",
                quantity,
                item,
                upc,
                f"Shipment {item}",
                "G26010101",
                "FACTORY-UNUSED",
                quantity // 2,
                quantity // 3,
                date(2026, 12, 1),
                date(2027, 1, 15),
                quantity / 6,
                quantity / 6 * 0.02,
            ]
        )
        cp.append(["TEST-PO-01", quantity, item, upc, f"Customer {item}", date(2026, 10, 1)])
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()


if __name__ == "__main__":
    generate(
        Path(sys.argv[1])
        if len(sys.argv) > 1
        else Path(__file__).with_name("synthetic_shefish.xlsx")
    )
