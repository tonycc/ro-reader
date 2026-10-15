"""PF options sheet 价格版本规则的行级选价测试。

规则基准（与真实 PF base 文件一致）：
- PI 行已发货且已签收（BALANCE QTY 有值、小于客户PO "Order Quantity"
  且 INV# 非空）→ 沿用上次发货的发票价：{seller}-INV 键 + 本行
  "ACTUAL EX FACTORY" floor 匹配断点。
- 其余 PI 行（BALANCE QTY 空、等于订单量、或已发货但 INV# 空的在途行）
  → {seller}-PI 键 + 客户PO "PO Creation Date" 匹配。
- Invoice 用 PO record "ACTUAL EX FACTORY" 在 {seller}-INV 键内匹配。
- 早于首断点/日期缺失 → 最早版本；晚于末断点 → 末版本。
- 版本标签与 DATA BASE 行1组标签严格字面匹配；组内取 COMBO 列。
"""

from __future__ import annotations

from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from ro_generator.document_model import (
    CODE_COMPONENT_PRICE_MISSING,
    CODE_LINE_NOT_PRICED,
    CODE_PRICE_VERSION_UNRESOLVED,
    build_pi_model,
)
from ro_generator.generator import PreviewResult, preview_from_snapshot
from ro_generator.models import DocumentRequest, DocumentType
from ro_generator.price_options import (
    PriceBook,
    build_price_book,
    build_price_groups,
)
from ro_generator.profiles import GenerationContext, create_pf_profile, profile_scope
from ro_generator.renderer import render_document
from ro_generator.resolver import ResolveResult, resolve_po_rows
from ro_generator.schema_inspect import inspect_schema
from ro_generator.template_mapping import load_template_mapping
from ro_generator.workbook_reader import SheetData, WorkbookReader
from ro_generator.workbook_snapshot import (
    BuildSnapshotError,
    WorkbookSnapshot,
    build_workbook_snapshot,
)

# —————————————————————————————————————
# 合成 fixture
# —————————————————————————————————————

# DATA BASE 布局：行1 组标签 / 行2 表头 / 行3 数据。
# - F/G：SK/YM 的两个版本组
# - H/I：EMAX 的两个版本组（DDP价格7/8 无列，覆盖"版本无组"路径）
# - J：中间商价格4（单列 COMBO，无 ROD/REEL，用于验证组件价缺失告警）
# - K..M：中间商价格5（COMBO + ROD + REEL，供组件拆分跟随组验证）
# - N/O：中间商价格6/7、厂家价格6/7（NEW PO 版 COMBO，即 data_base_price_columns
#        的固定列，用于验证 PO 取价不受 options 影响）
_DB_GROUP_LABELS = [
    None,
    None,
    None,
    None,
    None,
    "厂家价格4",
    "厂家价格5",
    "DDP价格5",
    "DDP价格6",
    "中间商价格4",
    "中间商价格5",
    None,
    None,
    "中间商价格6/7",
    "厂家价格6/7",
]
_DB_HEADERS = [
    "SAP",
    "Material Description",
    "Category",
    "round value",
    "MOQ",
    "GS-SK/YM COMBO FOB 2026",
    "GS-SK/YM COMBO FOB 20260612-OPO",
    "PF-EMAX COMBO DDP 2026 EFFECTIVE AS OF JUN/12/26-OPO",
    "PF-EMAX COMBO DDP 2026 EFFECTIVE AS OF JUN/12/26-NEW PO",
    "EMAX-GS COMBO FOB 2026 留下3%",
    "EMAX-GS COMBO FOB 20260612-OPO 留下3%",
    "EMAX-GS ROD FOB 20260612-OPO 留下3%",
    "EMAX-GS REEL FOB 20260612-OPO 留下3%",
    "EMAX-GS COMBO FOB 20260612-NEW PO 留下3%",
    "GS-SK/YM COMBO FOB 20260612-NEW PO",
]
_DB_VALUES = [
    "10001",
    "PF test item",
    "Single Rod",
    24,
    100,
    41,
    42,
    51,
    52,
    61,
    62,
    7,
    4,
    63,
    43,
]

_OPTIONS_HEADERS = [
    "EMAX-PI",
    "EMAX-PI取值",
    "GS-PI",
    "GS-PI取值",
    "SK/YM-PI",
    "SK/YM-PI取值",
    "EMAX-INV",
    "EMAX-INV取值",
    "GS-INV",
    "GS-INV取值字段",
    "SK/YM-INV",
    "SK/YM-INV取值",
]
# 断点结构与真实文件一致；EMAX-PI 的 0918 断点指向 fixture 中无对应组的
# 「DDP价格8」，覆盖"版本标签无匹配组"路径。
_OPTIONS_ROWS = [
    [
        20260612,
        "DDP价格6",
        20260612,
        "中间商价格4",
        20260612,
        "厂家价格4",
        20260612,
        "DDP价格5",
        20260612,
        "中间商价格4",
        20260612,
        "厂家价格4",
    ],
    [
        20260728,
        "DDP价格7",
        20260728,
        "中间商价格5",
        20260728,
        "厂家价格5",
        None,
        None,
        20260728,
        "中间商价格5",
        20260728,
        "厂家价格5",
    ],
    [20260918, "DDP价格8", None, None, None, None, None, None, None, None, None, None],
]


def _make_pf_base(
    tmp_path: Path,
    *,
    category: str = "Single Rod",
    po_creation_date: str = "2026-08-01",
    actual_ex_factory: str | None = "2026-09-15",
    balance_qty: object = None,
    inv_no: object = "G26010101",
    options_rows: list[list[object]] | None = None,
    include_options: bool = True,
    blank_price_columns: tuple[str, ...] = (),
    drop_options_value_headers: tuple[str, ...] = (),
) -> Path:
    workbook = Workbook()
    data_base = workbook.active
    assert data_base is not None
    data_base.title = "DATA BASE TEMPLATE"
    data_base.append(_DB_GROUP_LABELS)
    data_base.append(_DB_HEADERS)
    data_base_values: list[object] = list(_DB_VALUES)
    data_base_values[2] = category
    for header in blank_price_columns:
        data_base_values[_DB_HEADERS.index(header)] = None
    data_base.append(data_base_values)

    po_record = workbook.create_sheet("PO RECORD 26")
    headers: list[object] = ["PO NO.", "ITEM LINE#", "SAP Number", "INV#", 2601]
    if actual_ex_factory is not None or balance_qty is not None:
        headers += ["ACTUAL EX FACTORY", "BALANCE QTY"]
    po_record.append(headers)
    row: list[object] = ["4500000001", 10, "10001", inv_no, 90]
    if actual_ex_factory is not None or balance_qty is not None:
        row += [actual_ex_factory, balance_qty]
    po_record.append(row)

    customer_po = workbook.create_sheet("new PO template")
    customer_po.append(
        [
            "PO Creation Date",
            "PO#",
            "PO-Item",
            "Material",
            "Material Description",
            "PO requested ex-fty date",
            "Order Quantity",
        ]
    )
    customer_po.append(
        [po_creation_date, "4500000001", 10, "10001", "PF test item", "2026-09-01", 90]
    )

    if include_options:
        options = workbook.create_sheet("options")
        options_headers: list[object] = list(_OPTIONS_HEADERS)
        for key in drop_options_value_headers:
            # 清掉键列右邻的取值列表头：键仍在，但版本标签列失去表头
            options_headers[_OPTIONS_HEADERS.index(key) + 1] = None
        options.append(options_headers)
        for opt_row in options_rows if options_rows is not None else _OPTIONS_ROWS:
            options.append(opt_row)

    path = tmp_path / "pf-base.xlsx"
    workbook.save(path)
    return path


def _snapshot(base_file: Path) -> tuple[WorkbookSnapshot, GenerationContext]:
    profile = create_pf_profile()
    context = GenerationContext(profile=profile, base_file=base_file)
    return build_workbook_snapshot(str(base_file), context=context), context


def _resolve(snapshot: WorkbookSnapshot, context: GenerationContext) -> ResolveResult:
    return resolve_po_rows(
        snapshot.po_rows_for_po("4500000001"),
        snapshot.product_index,
        po_no="4500000001",
        customer_po_rows=snapshot.customer_po_rows_for_po("4500000001"),
        profile=context.profile,
        price_book=snapshot.price_book,
    )


def _preview(
    snapshot: WorkbookSnapshot,
    context: GenerationContext,
    seller: str,
    documents: tuple[DocumentType, ...],
) -> PreviewResult:
    request = DocumentRequest(
        base_file=str(context.base_path),
        po_no="4500000001",
        documents=documents,
        seller=seller,
        invoice_no="G26010101" if "INVOICE" in documents else None,
        output_dir=str(context.base_path.parent / "output"),
    )
    return preview_from_snapshot(snapshot, request, context=context)


# —————————————————————————————————————
# PI 订单行去重与缺价显示回归（仅合成 workbook）
# —————————————————————————————————————


def _make_split_pf_base(
    tmp_path: Path,
    *,
    completed: bool,
    category: str = "Single Rod",
) -> Path:
    path = _make_pf_base(tmp_path, category=category, po_creation_date="2026-01-01", balance_qty=50)
    workbook = load_workbook(path)
    po_record = workbook["PO RECORD 26"]
    po_record.delete_rows(2, po_record.max_row - 1)
    po_record.append(["4500000001", "00010", "10001", "G26010101", 40, "2026-07-01", 50])
    po_record.append(["4500000001", 10, "10001", "G26010102", 30, "2026-08-01", 20])
    if completed:
        po_record.append(["4500000001", 10, "10001", "G26010102", 20, "2026-09-15", 0])
    workbook.save(path)
    workbook.close()
    return path


@pytest.mark.parametrize("completed", [False, True], ids=["partial", "completed"])
@pytest.mark.parametrize(
    ("seller", "category", "price"),
    [
        ("GS PTE", "Single Rod", "62"),
        ("EMAX PTE", "Single Rod", "51"),
        ("YM", "Single Rod", "42"),
        ("SK", "Single Reel", "42"),
    ],
)
def test_pi_counts_split_customer_order_once(
    tmp_path: Path, completed: bool, seller: str, category: str, price: str
) -> None:
    snapshot, context = _snapshot(
        _make_split_pf_base(tmp_path, completed=completed, category=category)
    )
    resolved = _resolve(snapshot, context)
    assert len(resolved.lines) == (3 if completed else 2)  # 出货明细仍完整保留
    with profile_scope(context.profile):
        result = build_pi_model(
            resolved.lines,
            seller=seller,
            buyer=context.profile.rules.buyer_for(seller) or "",
            po_no="4500000001",
        )
    assert result.model is not None
    assert len(result.model.lines) == 1
    line = result.model.lines[0]
    assert line.quantity == result.model.total_quantity == Decimal("90")
    assert line.unit_price == Decimal(price)
    assert result.model.total_amount == Decimal(price) * 90
    assert line.source_row == (4 if completed else 3)
    assert line.price_source is not None
    assert line.price_source.option_key is not None
    assert line.price_source.option_key.endswith("-INV")
    preview = _preview(snapshot, context, seller, ("PI",))
    assert preview.preview is not None, preview.errors
    assert len(preview.preview.lines) == 1
    assert preview.preview.totals["total_quantity"] == "90"


def test_pi_does_not_merge_distinct_customer_items_with_same_sap(tmp_path: Path) -> None:
    path = _make_split_pf_base(tmp_path, completed=True)
    workbook = load_workbook(path)
    workbook["new PO template"].append(
        ["2026-01-01", "4500000001", 20, "10001", "Second item", "2026-09-01", 30]
    )
    workbook["PO RECORD 26"].append(["4500000001", "00020", "10001", None, None, None, 30])
    workbook.save(path)
    workbook.close()
    snapshot, context = _snapshot(path)
    preview = _preview(snapshot, context, "GS PTE", ("PI",))
    assert preview.preview is not None, preview.errors
    assert [line["quantity"] for line in preview.preview.lines] == ["90", "30"]
    assert [line["unit_price"] for line in preview.preview.lines] == ["$62.00", "$61.00"]
    assert preview.preview.totals["total_quantity"] == "120"
    assert preview.preview.totals["total_amount"] == "$7,410.00"


def test_pi_uses_latest_invoiced_shipment_not_last_physical_row(tmp_path: Path) -> None:
    path = _make_split_pf_base(tmp_path, completed=False)
    workbook = load_workbook(path)
    po_record = workbook["PO RECORD 26"]
    # 最新已开票行在上方；表尾另有更晚但未开票的在途行。
    first, second = list(po_record.values)[1:]
    for row_index, values in enumerate((second, first), start=2):
        for column_index, value in enumerate(values, start=1):
            po_record.cell(row_index, column_index).value = value
    po_record.append(["4500000001", 10, "10001", None, 20, "2026-09-15", 0])
    workbook.save(path)
    workbook.close()
    snapshot, context = _snapshot(path)
    with profile_scope(context.profile):
        result = build_pi_model(
            _resolve(snapshot, context).lines,
            seller="GS PTE",
            buyer="EMAX PTE",
            po_no="4500000001",
        )
    assert result.model is not None
    assert len(result.model.lines) == 1
    line = result.model.lines[0]
    assert line.source_row == 2
    assert line.unit_price == Decimal("62")
    assert line.price_source is not None
    assert line.price_source.line_date == date(2026, 8, 1)
    assert line.price_source.option_key == "GS-INV"


@pytest.mark.parametrize("document", ["INVOICE", "PL"])
def test_pi_grouping_does_not_collapse_invoice_shipments(
    tmp_path: Path, document: DocumentType
) -> None:
    snapshot, context = _snapshot(_make_split_pf_base(tmp_path, completed=True))
    request = DocumentRequest(
        base_file=str(context.base_path),
        po_no="4500000001",
        documents=(document,),
        seller="GS PTE",
        invoice_no="G26010102",
        output_dir=str(tmp_path / "output"),
    )
    result = preview_from_snapshot(snapshot, request, context=context)
    assert result.preview is not None, result.errors
    assert [line["quantity"] for line in result.preview.lines] == ["30", "20"]
    assert result.preview.totals["total_quantity"] == "50"


def test_pi_missing_one_price_does_not_show_partial_amount_as_total(tmp_path: Path) -> None:
    path = _make_pf_base(
        tmp_path,
        po_creation_date="2026-01-01",
        balance_qty=90,
        blank_price_columns=(_DB_HEADERS[10],),
    )
    workbook = load_workbook(path)
    workbook["new PO template"].append(
        ["2026-08-01", "4500000001", 20, "10001", "Unpriced item", "2026-09-01", 30]
    )
    workbook["PO RECORD 26"].append(["4500000001", 20, "10001", None, None, None, 30])
    workbook.save(path)
    workbook.close()
    snapshot, context = _snapshot(path)
    preview = _preview(snapshot, context, "GS PTE", ("PI",))
    assert preview.preview is not None, preview.errors
    assert [line["unit_price"] for line in preview.preview.lines] == ["$61.00", ""]
    assert [line["amount"] for line in preview.preview.lines] == ["$5,490.00", ""]
    assert preview.preview.totals["total_amount"] == ""
    assert preview.preview.totals["total_quantity"] == "120"
    assert any(m.code == CODE_LINE_NOT_PRICED and m.severity == "high" for m in preview.warnings)


@pytest.mark.parametrize("price", [None, 0], ids=["missing", "real-zero"])
def test_pi_missing_and_zero_price_remain_distinct_in_preview_and_excel(
    tmp_path: Path, price: int | None
) -> None:
    path = _make_pf_base(tmp_path)
    workbook = load_workbook(path)
    workbook["DATA BASE TEMPLATE"]["K3"] = price
    workbook.save(path)
    workbook.close()
    snapshot, context = _snapshot(path)
    preview = _preview(snapshot, context, "GS PTE", ("PI",))
    assert preview.preview is not None, preview.errors
    display = "" if price is None else "$0.00"
    assert preview.preview.lines[0]["unit_price"] == display
    assert preview.preview.lines[0]["amount"] == display
    assert preview.preview.totals["total_amount"] == display
    assert any(m.code == CODE_LINE_NOT_PRICED for m in preview.warnings) == (price is None)
    with profile_scope(context.profile):
        result = build_pi_model(
            _resolve(snapshot, context).lines,
            seller="GS PTE",
            buyer="EMAX PTE",
            po_no="4500000001",
        )
        assert result.model is not None
        mapping_path = context.profile.assets.mapping_path("GS PTE", "PI")
        assert mapping_path is not None
        mapping = load_template_mapping(mapping_path)
        rendered = render_document(result.model, mapping, tmp_path / "pi.xlsx")
    output = load_workbook(rendered.output_path)
    try:
        sheet = output[mapping.sheet]
        for key in ("unit_price", "amount"):
            cell = f"{getattr(mapping.lines.columns, key)}{mapping.lines.start_row}"
            assert sheet[cell].value == price
        for key in ("subtotal", "amount"):
            assert sheet[mapping.totals[key].cell].value == price
    finally:
        output.close()


# —————————————————————————————————————
# PriceBook 单元行为
# —————————————————————————————————————


def _pf_book(tmp_path: Path) -> PriceBook:
    base_file = _make_pf_base(tmp_path)
    profile = create_pf_profile()
    with WorkbookReader(str(base_file), schema=profile.schema) as reader:
        options_sheet = reader.read_sheet("options")
        group_labels = reader.read_group_labels("DATA BASE TEMPLATE")
        db_headers = reader.read_headers("DATA BASE TEMPLATE")
    config = profile.schema.price_options
    assert config is not None
    return build_price_book(
        options_sheet,
        group_labels,
        db_headers.header_columns,
        config,
        profile.schema,
    )


def test_price_groups_locate_combo_columns(tmp_path: Path) -> None:
    groups = _pf_book(tmp_path).groups
    assert groups["厂家价格4"].combo == "GS-SK/YM COMBO FOB 2026"
    assert groups["厂家价格5"].combo == "GS-SK/YM COMBO FOB 20260612-OPO"
    assert groups["中间商价格4"].combo == "EMAX-GS COMBO FOB 2026 留下3%"
    assert groups["中间商价格5"].combo == "EMAX-GS COMBO FOB 20260612-OPO 留下3%"
    assert groups["中间商价格5"].rod == "EMAX-GS ROD FOB 20260612-OPO 留下3%"
    assert groups["中间商价格5"].reel == "EMAX-GS REEL FOB 20260612-OPO 留下3%"
    ddp5_combo = groups["DDP价格5"].combo
    assert ddp5_combo is not None and ddp5_combo.startswith("PF-EMAX COMBO DDP")
    assert "DDP价格8" not in groups


def test_breakpoint_floor_matching(tmp_path: Path) -> None:
    book = _pf_book(tmp_path)
    # 恰好等于断点
    assert book.resolve("PI", "EMAX PTE", date(2026, 7, 28)).version_label == "DDP价格7"
    # 两断点之间 → floor
    assert book.resolve("PI", "EMAX PTE", date(2026, 8, 1)).version_label == "DDP价格7"
    # 早于首断点 → 最早版本
    assert book.resolve("PI", "EMAX PTE", date(2026, 1, 1)).version_label == "DDP价格6"
    # 晚于末断点 → 末版本
    hit = book.resolve("PI", "EMAX PTE", date(2027, 1, 1))
    assert hit.version_label == "DDP价格8"
    assert hit.kind == "missing"  # fixture 中无 DDP价格8 组
    # 日期缺失 → 最早版本
    hit = book.resolve("PI", "GS PTE", None)
    assert hit.version_label == "中间商价格4"
    assert hit.column == "EMAX-GS COMBO FOB 2026 留下3%"


def test_strict_label_matching_no_fuzzy(tmp_path: Path) -> None:
    """「厂家价格4」≠「厂家价格3/4」：字面不等即不命中。"""
    book = _pf_book(tmp_path)
    hit = book.resolve("PI", "SK", date(2026, 6, 12))
    assert hit.kind == "options"
    assert hit.column == "GS-SK/YM COMBO FOB 2026"
    # 组标签改为「厂家价格3/4」时，「厂家价格4」不得模糊命中
    groups = build_price_groups({6: "厂家价格3/4"}, {"GS-SK/YM COMBO FOB 2026": 6})
    book2 = PriceBook(
        options=book.options,
        groups=groups,
        seller_keys=book.seller_keys,
        doc_keys=book.doc_keys,
    )
    hit2 = book2.resolve("PI", "SK", date(2026, 6, 12))
    assert hit2.kind == "missing"
    assert hit2.version_label == "厂家价格4"


def test_options_missing_key_column_is_issue(tmp_path: Path) -> None:
    base_file = _make_pf_base(tmp_path)
    profile = create_pf_profile()
    with WorkbookReader(str(base_file), schema=profile.schema) as reader:
        options_sheet = reader.read_sheet("options")
        shrunk = SheetData(
            sheet_name=options_sheet.sheet_name,
            headers=tuple(h for h in options_sheet.headers if h != "EMAX-PI"),
            header_columns={
                h: c for h, c in options_sheet.header_columns.items() if h != "EMAX-PI"
            },
            rows=options_sheet.rows,
        )
        group_labels = reader.read_group_labels("DATA BASE TEMPLATE")
        db_headers = reader.read_headers("DATA BASE TEMPLATE")
    config = profile.schema.price_options
    assert config is not None
    book = build_price_book(
        shrunk,
        group_labels,
        db_headers.header_columns,
        config,
        profile.schema,
    )
    assert any("EMAX-PI" in issue for issue in book.issues)
    assert book.resolve("PI", "EMAX PTE", date(2026, 8, 1)).kind == "missing"


# —————————————————————————————————————
# 行级解析
# —————————————————————————————————————


def test_pi_unit_price_uses_po_creation_date_version(tmp_path: Path) -> None:
    """PI：PO Creation Date 2026-08-01 → GS-PI 命中 0728 断点 → 中间商价格5。"""
    snapshot, context = _snapshot(_make_pf_base(tmp_path))
    line = _resolve(snapshot, context).lines[0]
    assert line.price_sources[("PI", "GS PTE")].version_label == "中间商价格5"
    assert line.pi_prices[("GS PTE", "EMAX PTE")] == Decimal("62")
    # 静态固定列不被 PI 版本价污染（PO 取价来源）
    assert line.prices[("GS PTE", "EMAX PTE")] == Decimal("63")

    preview = _preview(snapshot, context, "GS PTE", ("PI",))
    assert preview.status == "success", preview.errors
    assert preview.preview is not None
    assert preview.preview.lines[0]["unit_price"] == "$62.00"
    source = {e["preview_field"]: e for e in preview.preview.source_entries}["line[0].unit_price"]
    assert source["field"] == "EMAX-GS COMBO FOB 20260612-OPO 留下3%"
    assert "中间商价格5" in str(source["rule"])


def test_pi_unit_price_before_first_breakpoint_uses_earliest(tmp_path: Path) -> None:
    snapshot, context = _snapshot(_make_pf_base(tmp_path, po_creation_date="2026-01-01"))
    line = _resolve(snapshot, context).lines[0]
    assert line.price_sources[("PI", "SK")].version_label == "厂家价格4"
    assert line.pi_prices[("SK", "YM")] == Decimal("41")
    assert line.prices[("SK", "YM")] == Decimal("43")


def test_invoice_unit_price_uses_actual_ex_factory(tmp_path: Path) -> None:
    """Invoice：ACTUAL EX FACTORY 2026-09-15 → GS-INV 命中 0728 → 中间商价格5。"""
    snapshot, context = _snapshot(_make_pf_base(tmp_path))
    line = _resolve(snapshot, context).lines[0]
    assert line.price_sources[("INVOICE", "GS PTE")].version_label == "中间商价格5"
    assert line.invoice_prices[("GS PTE", "EMAX PTE")] == Decimal("62")

    preview = _preview(snapshot, context, "GS PTE", ("INVOICE",))
    assert preview.status == "success", preview.errors
    assert preview.preview is not None
    assert preview.preview.lines[0]["unit_price"] == "$62.00"


def test_invoice_before_first_breakpoint_uses_earliest(tmp_path: Path) -> None:
    snapshot, context = _snapshot(_make_pf_base(tmp_path, actual_ex_factory="2026-01-01"))
    line = _resolve(snapshot, context).lines[0]
    assert line.price_sources[("INVOICE", "EMAX PTE")].version_label == "DDP价格5"
    assert line.invoice_prices[("EMAX PTE", "PF")] == Decimal("51")


def test_pi_fully_shipped_row_uses_invoice_price(tmp_path: Path) -> None:
    """BALANCE QTY=0 已发完且 INV# 非空：沿用 -INV 键 + ACTUAL EX FACTORY。

    PO Creation Date 2026-01-01（若走 -PI 键会命中最早版本 中间商价格4=61），
    ACTUAL EX FACTORY 2026-09-15 → GS-INV 命中 0728 断点 → 中间商价格5=62。
    """
    snapshot, context = _snapshot(
        _make_pf_base(tmp_path, balance_qty=0, po_creation_date="2026-01-01")
    )
    line = _resolve(snapshot, context).lines[0]
    resolved = line.price_sources[("PI", "GS PTE")]
    assert resolved.kind == "options"
    assert resolved.option_key == "GS-INV"
    assert resolved.line_date == date(2026, 9, 15)
    assert resolved.version_label == "中间商价格5"
    assert line.pi_prices[("GS PTE", "EMAX PTE")] == Decimal("62")
    # 静态固定列值留在 prices 供 PO，不受 PI 版本价影响
    assert line.prices[("GS PTE", "EMAX PTE")] == Decimal("63")
    # Invoice 取价与 PI 分支无关
    assert line.invoice_prices[("GS PTE", "EMAX PTE")] == Decimal("62")

    preview = _preview(snapshot, context, "GS PTE", ("PI",))
    assert preview.status == "success", preview.errors
    assert preview.preview is not None
    assert preview.preview.lines[0]["unit_price"] == "$62.00"
    source = {e["preview_field"]: e for e in preview.preview.source_entries}["line[0].unit_price"]
    assert source["field"] == "EMAX-GS COMBO FOB 20260612-OPO 留下3%"
    assert "GS-INV" in str(source["rule"])


def test_pi_partial_shipment_with_invoice_uses_inv_key(tmp_path: Path) -> None:
    """BALANCE QTY < Order Quantity 且 INV# 非空（部分发货已签收）→ -INV 键。"""
    snapshot, context = _snapshot(
        _make_pf_base(tmp_path, balance_qty=30, po_creation_date="2026-01-01")
    )
    line = _resolve(snapshot, context).lines[0]
    resolved = line.price_sources[("PI", "GS PTE")]
    assert resolved.kind == "options"
    assert resolved.option_key == "GS-INV"
    assert resolved.version_label == "中间商价格5"
    assert line.pi_prices[("GS PTE", "EMAX PTE")] == Decimal("62")


def test_pi_partial_shipment_in_transit_uses_pi_key(tmp_path: Path) -> None:
    """BALANCE QTY < Order Quantity 但 INV# 为空（在途未签收）→ -PI 键。"""
    snapshot, context = _snapshot(
        _make_pf_base(tmp_path, balance_qty=30, inv_no=None, po_creation_date="2026-01-01")
    )
    line = _resolve(snapshot, context).lines[0]
    resolved = line.price_sources[("PI", "GS PTE")]
    assert resolved.kind == "options"
    assert resolved.option_key == "GS-PI"
    assert resolved.line_date == date(2026, 1, 1)
    assert resolved.version_label == "中间商价格4"
    assert line.pi_prices[("GS PTE", "EMAX PTE")] == Decimal("61")


def test_pi_unshipped_balance_equals_order_qty_uses_pi_key(tmp_path: Path) -> None:
    """BALANCE QTY = 客户PO Order Quantity（未发货）→ -PI 键。"""
    snapshot, context = _snapshot(
        _make_pf_base(tmp_path, balance_qty=90, po_creation_date="2026-01-01")
    )
    line = _resolve(snapshot, context).lines[0]
    resolved = line.price_sources[("PI", "GS PTE")]
    assert resolved.kind == "options"
    assert resolved.option_key == "GS-PI"
    assert resolved.version_label == "中间商价格4"
    assert line.pi_prices[("GS PTE", "EMAX PTE")] == Decimal("61")


def test_missing_version_group_warns_and_price_empty(tmp_path: Path) -> None:
    """EMAX-PI 命中 DDP价格8（无对应组）→ 装配层告警 + PI 单价为空（不回退静态列）。

    版本解析失败属于单据级上下文：resolver 只记录 ResolvedPrice，告警在
    装配 EMAX PI 时才产生，不污染其他卖方/单据的预览。
    """
    snapshot, context = _snapshot(_make_pf_base(tmp_path, po_creation_date="2026-10-01"))
    result = _resolve(snapshot, context)
    line = result.lines[0]
    resolved = line.price_sources[("PI", "EMAX PTE")]
    assert resolved.kind == "missing"
    assert resolved.version_label == "DDP价格8"
    # resolver 不再在解析阶段产生版本告警
    assert not any(m.code == CODE_PRICE_VERSION_UNRESOLVED for m in result.messages)

    preview = _preview(snapshot, context, "EMAX PTE", ("PI",))
    assert preview.status == "success", preview.errors
    assert preview.preview is not None
    # 缺价不能伪装成零价，也不能给出完整金额。
    assert preview.preview.lines[0]["unit_price"] == ""
    assert preview.preview.lines[0]["amount"] == ""
    assert preview.preview.totals["total_amount"] == ""
    warning_codes = {m.code for m in preview.warnings}
    assert "LINE_NOT_PRICED_FOR_SEGMENT" in warning_codes
    assert CODE_PRICE_VERSION_UNRESOLVED in warning_codes
    version_warnings = [m for m in preview.warnings if m.code == CODE_PRICE_VERSION_UNRESOLVED]
    assert version_warnings and "DDP价格8" in version_warnings[0].message
    # pi_prices 不留任何该段的版本价，静态列旧值也不被搬走
    assert ("EMAX PTE", "PF") not in line.pi_prices
    assert line.prices[("EMAX PTE", "PF")] == Decimal("52")


def test_missing_options_sheet_blocks_snapshot(tmp_path: Path) -> None:
    base_file = _make_pf_base(tmp_path, include_options=False)
    profile = create_pf_profile()
    context = GenerationContext(profile=profile, base_file=base_file)
    with pytest.raises(BuildSnapshotError, match="options"):
        build_workbook_snapshot(str(base_file), context=context)


def test_missing_options_key_header_blocks_snapshot(tmp_path: Path) -> None:
    base_file = _make_pf_base(tmp_path)
    workbook = load_workbook(base_file)
    options = workbook["options"]
    options["A1"] = "EMAX-PI-RENAMED"
    workbook.save(base_file)
    workbook.close()
    profile = create_pf_profile()
    context = GenerationContext(profile=profile, base_file=base_file)
    with pytest.raises(BuildSnapshotError, match="EMAX-PI"):
        build_workbook_snapshot(str(base_file), context=context)


def test_invoice_component_prices_follow_selected_group(tmp_path: Path) -> None:
    """GS Invoice 组件拆分列跟随所选版本组（组内 ROD/REEL 列）。

    ACTUAL EX FACTORY 2026-09-15 → GS-INV 命中 中间商价格5 组，
    ROD/REEL 组件价取组内 OPO 列（7/4），而非静态 NEW PO 配置列。
    """
    snapshot, context = _snapshot(_make_pf_base(tmp_path, category="Combo"))
    line = _resolve(snapshot, context).lines[0]
    assert line.resolved_component_prices["GS PTE/rod"] == Decimal("7")
    assert line.resolved_component_prices["GS PTE/reel"] == Decimal("4")

    preview = _preview(snapshot, context, "GS PTE", ("INVOICE",))
    assert preview.status == "success", preview.errors
    assert preview.preview is not None
    breakdown = preview.preview.cost_breakdown
    assert [item["component"] for item in breakdown] == ["RODS", "REELS"]
    assert breakdown[0]["unit_price"] == "$7.00"
    assert breakdown[1]["unit_price"] == "$4.00"


# —————————————————————————————————————
# 评审回归（6 项）
# —————————————————————————————————————


def test_pi_version_price_does_not_leak_into_po(tmp_path: Path) -> None:
    """PO 始终读静态固定列：SK/YM-PI 命中的版本价不得改变 GS PTE 的 PO 单价。

    回归：PI 版本价曾回写 line.prices，GS PTE 的 PO 输出 $42（SK/YM-PI 的
    OPO 版本价）而非静态 NEW PO 列的 $43。
    """
    snapshot, context = _snapshot(_make_pf_base(tmp_path))
    line = _resolve(snapshot, context).lines[0]
    # SK/YM-PI 命中 0728 断点 → 厂家价格5 → OPO 列 42；与 PO 段键相同
    assert line.pi_prices[("YM", "GS PTE")] == Decimal("42")
    # 同段键的静态列值（NEW PO 列）必须原样保留
    assert line.prices[("YM", "GS PTE")] == Decimal("43")

    preview = _preview(snapshot, context, "GS PTE", ("PO",))
    assert preview.status == "success", preview.errors
    assert preview.preview is not None
    assert preview.preview.lines[0]["unit_price"] == "$43.00"


def test_blank_selected_version_cell_does_not_fall_back(tmp_path: Path) -> None:
    """选中版本列但该 SAP 行单元格为空 → 缺价告警，不回退静态列旧价。

    回归：pi_prices 曾预填静态列值，选中列取空时静默沿用旧价 $63，
    而来源摘要却显示版本列。
    """
    snapshot, context = _snapshot(
        _make_pf_base(
            tmp_path,
            blank_price_columns=("EMAX-GS COMBO FOB 20260612-OPO 留下3%",),
        )
    )
    line = _resolve(snapshot, context).lines[0]
    resolved = line.price_sources[("PI", "GS PTE")]
    assert resolved.kind == "options"
    assert resolved.column == "EMAX-GS COMBO FOB 20260612-OPO 留下3%"
    # 选中列取空：pi_prices 不留值；静态列旧值留在 prices 供 PO，不被 PI 使用
    assert ("GS PTE", "EMAX PTE") not in line.pi_prices
    assert line.prices[("GS PTE", "EMAX PTE")] == Decimal("63")

    preview = _preview(snapshot, context, "GS PTE", ("PI",))
    assert preview.status == "success", preview.errors
    assert preview.preview is not None
    assert preview.preview.lines[0]["unit_price"] == ""
    assert preview.preview.lines[0]["amount"] == ""
    assert preview.preview.totals["total_amount"] == ""
    assert "LINE_NOT_PRICED_FOR_SEGMENT" in {m.code for m in preview.warnings}


def test_price_version_warnings_scoped_to_document_and_seller(tmp_path: Path) -> None:
    """价格版本告警只出现在对应单据+卖方的装配结果里。

    回归：resolver 曾一次性产出所有卖方×单据族的告警，导致 GS PO 预览
    收到无关的 EMAX PI 版本缺失告警。
    """
    snapshot, context = _snapshot(_make_pf_base(tmp_path, po_creation_date="2026-10-01"))

    emax_pi = _preview(snapshot, context, "EMAX PTE", ("PI",))
    assert CODE_PRICE_VERSION_UNRESOLVED in {m.code for m in emax_pi.warnings}

    cases: tuple[tuple[str, tuple[DocumentType, ...]], ...] = (
        ("GS PTE", ("PI",)),
        ("GS PTE", ("PO",)),
        ("SK", ("PI",)),
    )
    for seller, documents in cases:
        preview = _preview(snapshot, context, seller, documents)
        codes = {m.code for m in preview.warnings}
        assert CODE_PRICE_VERSION_UNRESOLVED not in codes, (seller, documents)


def test_options_key_issue_is_not_remap_options(tmp_path: Path) -> None:
    """options 键列缺失标记为不可重映射。

    override 机制不支持 options 键列，修复向导不得提供保存后静默失效的列映射。
    """
    base_file = _make_pf_base(tmp_path)
    workbook = load_workbook(base_file)
    workbook["options"]["A1"] = "EMAX-PI-RENAMED"
    workbook.save(base_file)
    workbook.close()

    profile = create_pf_profile()
    with WorkbookReader(str(base_file), schema=profile.schema) as reader:
        inspection = inspect_schema(reader, profile.schema)
    options_issues = [
        issue for issue in inspection.field_issues if issue.logical_sheet == "options"
    ]
    assert options_issues, "应报告 EMAX-PI 键列缺失"
    assert any(issue.expected_header == "EMAX-PI" for issue in options_issues)
    assert all(issue.remappable is False for issue in options_issues)
    assert all(issue.repair_hint for issue in options_issues)


def test_missing_component_price_warns_on_combo_invoice(tmp_path: Path) -> None:
    """Combo Invoice 所选价格组缺 ROD/REEL 列 → 高优先级告警，不静默丢拆分。

    回归：组件价缺失时 cost_breakdown 曾直接为空且无告警，Invoice 显示
    success 却少了 RODS/REELS 两行。
    """
    snapshot, context = _snapshot(
        _make_pf_base(tmp_path, category="Combo", actual_ex_factory="2026-06-20")
    )
    line = _resolve(snapshot, context).lines[0]
    # 0612 断点 → 中间商价格4：组内只有 COMBO 列，无 ROD/REEL
    assert line.price_sources[("INVOICE", "GS PTE")].version_label == "中间商价格4"
    assert "GS PTE/rod" not in line.resolved_component_prices
    assert "GS PTE/reel" not in line.resolved_component_prices

    preview = _preview(snapshot, context, "GS PTE", ("INVOICE",))
    assert preview.status == "success", preview.errors
    assert preview.preview is not None
    # 主单价照常（COMBO 列有价），拆分区为空但必须有高优先级告警
    assert preview.preview.lines[0]["unit_price"] == "$61.00"
    assert not preview.preview.cost_breakdown
    component_warnings = [m for m in preview.warnings if m.code == CODE_COMPONENT_PRICE_MISSING]
    assert component_warnings, "缺组件价必须有告警"
    assert all(m.severity == "high" for m in component_warnings)
    assert any("RODS" in m.message for m in component_warnings)
    assert any("REELS" in m.message for m in component_warnings)


def test_missing_options_value_header_blocks_snapshot(tmp_path: Path) -> None:
    """options 键列右邻取值列表头缺失 → 结构校验阻断 snapshot。

    回归：此前该问题只进 PriceBook issues（低危告警），snapshot 成功后
    版本标签读不到，PI 静默输出 $0.00。
    """
    base_file = _make_pf_base(tmp_path, drop_options_value_headers=("EMAX-PI",))
    profile = create_pf_profile()
    context = GenerationContext(profile=profile, base_file=base_file)
    with pytest.raises(BuildSnapshotError, match="EMAX-PI"):
        build_workbook_snapshot(str(base_file), context=context)
