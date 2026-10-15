"""SHEFISH 结构、价格、组件采购和出货单据的合成回归。"""

from __future__ import annotations

import subprocess
import sys
from dataclasses import replace
from datetime import date
from pathlib import Path
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook
from ro_generator.base_schema import SchemaOverride
from ro_generator.generator import generate, preview_from_snapshot
from ro_generator.invoice_inspection import inspect_invoice_group_from_snapshot
from ro_generator.models import DocumentRequest, DocumentType
from ro_generator.profiles import GenerationContext, create_shefish_profile
from ro_generator.template_mapping import load_template_mapping
from ro_generator.workbench_service import ExportDocumentGroup, export_document_groups
from ro_generator.workbook_snapshot import build_workbook_snapshot


@pytest.fixture
def context(tmp_path: Path) -> GenerationContext:
    path = tmp_path / "shefish.xlsx"
    script = Path(__file__).resolve().parents[3] / "tests/fixtures/generate_shefish_base.py"
    subprocess.run([sys.executable, str(script), str(path)], check=True)
    return GenerationContext(create_shefish_profile(), path)


def test_profile_assets_are_isolated_and_sanitized(context: GenerationContext) -> None:
    profile = context.profile
    paths = list(profile.assets.root.glob("templates/*/mappings/*.yaml"))
    assert len(paths) == 12
    assert profile.rules.buyer_for("SK") == "GS PTE"
    assert profile.rules.buyer_for("EMAX PTE") == "SHEFISH"
    assert not profile.capabilities.documents_for("YM")
    for path in paths:
        mapping = load_template_mapping(path)
        assert mapping.preview_static_values["title"]
        assert mapping.preview_static_values["seller_info"]
        assert mapping.template_path.is_relative_to(profile.assets.root)
        with ZipFile(mapping.template_path) as archive:
            assert not any("externalLinks" in name for name in archive.namelist())
        workbook = load_workbook(mapping.template_path)
        for row in workbook[mapping.sheet]:
            for cell in row:
                assert cell.data_type != "f"
                assert cell.comment is None
                assert not any(
                    text in str(cell.value)
                    for text in ("DATA BASE", "固定值", "new PO", "PO RECORD", "生成日期")
                )
        workbook.close()


def test_snapshot_uses_item_number_without_sap_or_item_line(context: GenerationContext) -> None:
    snapshot = build_workbook_snapshot(str(context.base_file), context=context)
    assert len(snapshot.product_index) == 3
    assert snapshot.po_summary[0].status == "ready"
    assert set(snapshot.po_summary[0].sellers) == {"GS PTE", "EMAX PTE", "SK"}
    assert snapshot.invoice_summary[0].seller_invoice_numbers == {
        "GS PTE": "G26010101",
        "EMAX PTE": "G26010101",
        "SK": "G26010101",
    }
    inspection = inspect_invoice_group_from_snapshot(
        snapshot, snapshot.invoice_summary[0].invoice_group_key, context=context
    )
    assert inspection.rows[0].sellers == ("GS PTE", "EMAX PTE", "SK")
    assert inspection.rows[1].sellers == ("GS PTE", "EMAX PTE")


@pytest.mark.parametrize(
    "seller,document,prices,quantities",
    [
        ("EMAX PTE", "PI", ["$20.00", "$8.00", "$12.00"], ["120", "60", "90"]),
        ("EMAX PTE", "PO", ["$12.00", "$5.00", "$7.00"], ["120", "60", "90"]),
        ("GS PTE", "PI", ["$12.00", "$5.00", "$7.00"], ["120", "60", "90"]),
        ("SK", "PI", ["$6.00", "$6.00"], ["120", "90"]),
        ("EMAX PTE", "INVOICE", ["$20.00", "$8.00", "$12.00"], ["60", "30", "45"]),
        ("SK", "INVOICE", ["$6.00", "$6.00"], ["60", "45"]),
    ],
)
def test_prices_quantity_upc_and_sources(
    context: GenerationContext,
    seller: str,
    document: DocumentType,
    prices: list[str],
    quantities: list[str],
) -> None:
    snapshot = build_workbook_snapshot(str(context.base_file), context=context)
    result = preview_from_snapshot(
        snapshot,
        DocumentRequest(
            str(context.base_file),
            "TEST-PO-01",
            (document,),
            seller=seller,
        ),
        context=context,
    )
    assert result.status == "success", result.errors
    preview = result.preview
    assert preview is not None
    assert [line["unit_price"] for line in preview.lines] == prices
    assert [line["quantity"] for line in preview.lines] == quantities
    assert preview.lines[0]["upc"] == "000123456789"
    if document in {"PI", "PO"}:
        assert preview.lines[0]["description"] == "Customer TEST-COMBO"
        assert preview.resolved_values["document_date"] == "2026-10-01"
        assert preview.lines[0]["confirmed_ex_factory_date"] == date(2026, 12, 1)
    else:
        assert preview.lines[0]["description"] == "Shipment TEST-COMBO"
        assert preview.resolved_values["eta"] == "2027-01-15"
        assert preview.resolved_values["document_date"] == date.today().isoformat()
        assert any(entry["field"] == "2601" for entry in preview.source_entries)
    if seller == "EMAX PTE" and document == "PI":
        assert preview.pi_no == "TEST-PI-01"


def test_gs_po_previews_and_exports_both_component_factories(
    context: GenerationContext, tmp_path: Path
) -> None:
    snapshot = build_workbook_snapshot(str(context.base_file), context=context)
    request = DocumentRequest(
        str(context.base_file),
        "TEST-PO-01",
        ("PO",),
        seller="GS PTE",
        output_dir=str(tmp_path / "out"),
    )
    result = preview_from_snapshot(snapshot, request, context=context)
    assert result.status == "success", result.errors
    assert [p.title for p in result.previews] == ["PURCHASE ORDER · SK", "PURCHASE ORDER · YM"]
    sk, ym = result.previews
    assert [line["item_number"] for line in sk.lines] == ["TEST-COMBO", "TEST-REEL"]
    assert [line["item_number"] for line in ym.lines] == ["TEST-COMBO", "TEST-ROD"]
    assert sk.lines[0]["unit_price"] == "$6.00"
    assert ym.lines[0]["unit_price"] == "$4.00"
    assert any(entry["field"] == "YM ROD FOB 2026" for entry in ym.source_entries)
    exported = generate(request, context=context)
    assert exported.status == "success", exported.errors
    assert exported.output_file is not None and exported.output_file.endswith(".zip")
    with ZipFile(exported.output_file) as archive:
        assert len(archive.namelist()) == 2
        for filename in archive.namelist():
            with archive.open(filename) as stream:
                wb = load_workbook(stream)
                sheet = wb.active
                assert sheet is not None
                is_sk = filename.endswith("_SK.xlsx")
                assert sheet["E18"].value == (6 if is_sk else 4)
                assert sheet["G22"].value == (1260 if is_sk else 720)
                assert sheet["C18"].value == "000123456789"
                assert ("GUANGDONG" if is_sk else "WEIHAI") in sheet["F13"].value
                wb.close()


def test_missing_emax_pi_number_is_not_generated(context: GenerationContext) -> None:
    wb = load_workbook(context.base_file)
    for row in range(2, 5):
        wb["PO RECORD 2026"].cell(row, 2).value = None
    wb.save(context.base_file)
    wb.close()
    result = generate(
        DocumentRequest(str(context.base_file), "TEST-PO-01", ("PI",), seller="EMAX PTE"),
        context=context,
    )
    assert result.status == "error"
    assert result.errors[0].code == "PI_NO_MISSING"


def test_customer_po_can_precede_record_but_cannot_invent_emax_pi(
    context: GenerationContext,
) -> None:
    wb = load_workbook(context.base_file)
    wb["PO RECORD 2026"].delete_rows(2, 3)
    wb.save(context.base_file)
    wb.close()
    snapshot = build_workbook_snapshot(str(context.base_file), context=context)
    assert not snapshot.invoice_summary
    result = preview_from_snapshot(
        snapshot,
        DocumentRequest(str(context.base_file), "TEST-PO-01", ("PO",), seller="GS PTE"),
        context=context,
    )
    assert result.status == "success", result.errors
    assert len(result.previews) == 2
    assert result.previews[0].lines[0]["quantity"] == "120"
    missing_pi = preview_from_snapshot(
        snapshot,
        DocumentRequest(str(context.base_file), "TEST-PO-01", ("PI",), seller="EMAX PTE"),
        context=context,
    )
    assert missing_pi.errors[0].code == "PI_NO_MISSING"


def test_gs_variants_honor_workspace_price_column_override(context: GenerationContext) -> None:
    schema = context.schema.with_override(
        SchemaOverride(price_columns={"data_base_price_columns": {"YM/combo": "NEW ROD PRICE"}})
    )
    context = replace(context, profile=replace(context.profile, schema=schema))
    wb = load_workbook(context.base_file)
    column = wb["DATA BASE"].max_column + 1
    wb["DATA BASE"].cell(1, column, "NEW ROD PRICE")
    wb["DATA BASE"].cell(2, column, 8)
    wb.save(context.base_file)
    wb.close()
    snapshot = build_workbook_snapshot(str(context.base_file), context=context)
    result = preview_from_snapshot(
        snapshot,
        DocumentRequest(str(context.base_file), "TEST-PO-01", ("PO",), seller="GS PTE"),
        context=context,
    )
    assert result.previews[1].lines[0]["unit_price"] == "$8.00"
    assert any(entry["field"] == "NEW ROD PRICE" for entry in result.previews[1].source_entries)


def test_batch_export_contains_documents_not_nested_archives(
    context: GenerationContext, tmp_path: Path
) -> None:
    result = export_document_groups(
        base_file=str(context.base_file),
        po_no="TEST-PO-01",
        output_dir=str(tmp_path / "batch"),
        groups=(
            ExportDocumentGroup("GS PTE", ("PO",)),
            ExportDocumentGroup("EMAX PTE", ("PI",)),
        ),
        context=context,
    )
    assert result.status == "success", result.errors
    assert result.output_file is not None
    with ZipFile(result.output_file) as archive:
        assert len(archive.namelist()) == 3
        assert all(name.endswith(".xlsx") for name in archive.namelist())


def test_gs_po_inserts_rows_without_overwriting_totals(
    context: GenerationContext, tmp_path: Path
) -> None:
    wb = load_workbook(context.base_file)
    for index in range(4):
        item = f"ADDED-COMBO-{index}"
        for sheet_name, column in (("DATA BASE", 1), ("PO RECORD 2026", 5), ("New PO", 3)):
            row = [cell.value for cell in wb[sheet_name][2]]
            row[column - 1] = item
            wb[sheet_name].append(row)
    wb.save(context.base_file)
    wb.close()
    result = generate(
        DocumentRequest(
            str(context.base_file),
            "TEST-PO-01",
            ("PO",),
            seller="GS PTE",
            output_dir=str(tmp_path / "insert"),
        ),
        context=context,
    )
    assert result.status == "success", result.errors
    assert result.output_file is not None
    with ZipFile(result.output_file) as archive:
        for name in archive.namelist():
            with archive.open(name) as source:
                wb = load_workbook(source)
                sheet = wb.active
                assert sheet is not None
                assert sheet["B17"].value == "Item Number"
                assert sheet["B23"].value == "ADDED-COMBO-3"
                assert sheet["G24"].value == (4140 if name.endswith("_SK.xlsx") else 2640)
                assert sheet["F27"].value == "Cindy"
                wb.close()


def test_invoice_month_does_not_fall_back_to_another_shipment(context: GenerationContext) -> None:
    rules, schema = context.rules, context.schema
    assert rules.shipment_quantity_for_row({"INV#": "G26030101", "2601": 100}, schema) == (
        None,
        "2603",
    )
    assert rules.shipment_quantity_for_row({"INV#": "G27010101", "2701": 24}, schema) == (
        24,
        "2701",
    )


@pytest.mark.parametrize("seller", ["GS PTE", "EMAX PTE", "SK"])
def test_invoice_and_pl_are_separate_templates(
    context: GenerationContext, tmp_path: Path, seller: str
) -> None:
    result = generate(
        DocumentRequest(
            str(context.base_file),
            "TEST-PO-01",
            ("INVOICE", "PL"),
            seller=seller,
            output_format="zip",
            output_dir=str(tmp_path / seller),
        ),
        context=context,
    )
    assert result.status == "success", result.errors
    assert len(result.files) == 2
    assert {message.code for message in result.warnings} <= {"FORMULA_FALLBACK"}
    assert result.output_file is not None
    with ZipFile(result.output_file) as archive:
        pl_name = next(name for name in archive.namelist() if "-PL-" in name)
        with archive.open(pl_name) as stream:
            wb = load_workbook(stream)
            sheet = wb["PL"]
            row = 22 if seller == "GS PTE" else 23
            assert sheet[f"F{row}"].value == 60
            assert sheet[f"G{row}"].value == 10
            assert sheet[f"H{row}"].value == 20
            assert sheet[f"I{row}"].value == 30
            assert sheet[f"J{row}"].value == 0.2
            assert sheet[f"A{row}"].value == 1
            assert sheet[f"B{row}"].value == 10
            wb.close()
